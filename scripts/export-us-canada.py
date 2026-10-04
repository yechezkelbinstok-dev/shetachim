import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
d=json.load(open('web/data/centers.geojson'))['features']
sh={s['id']:s for s in json.load(open('web/data/shetachim.json'))['shetachim']}
def head(h):
    if isinstance(h,list): return ', '.join(head(x) for x in h)
    if isinstance(h,dict): return h.get('name') or json.dumps(h)
    return h or ''
FEM=re.compile(r'^(mrs|ms|miss|rebbetzin|rebbitzen|mme|madame|sra|frau)\.?$',re.I)
male=lambda p: not FEM.match((p.get('title') or '').strip())
nm=lambda p: ' '.join(x for x in [p.get('firstName'),p.get('lastName')] if x)
def main(ps):
    living=[p for p in ps if not p.get('isDeceased')]; passed=[p for p in ps if p.get('isDeceased')]
    pr=next((p for p in living if male(p)),None)
    w=None if pr else next((p for p in living if not male(p) and p.get('lastName')),None)
    wid=w and any(male(p) for p in passed)
    if pr: m=nm(pr)
    elif wid: m=f"{w['lastName']} family"
    elif w: m=nm(w)
    else: m=''
    others=[p for p in living if p is not pr and not (w and not wid and p is w)]
    return m, '; '.join(nm(p) for p in others), '; '.join(nm(p)+' obm' for p in passed)
C=[];P=[]
for f in d:
    pr=f['properties']
    if pr['country'] not in ('US','CA'): continue
    lon,lat=f['geometry']['coordinates']
    sid=pr['piece'].split(':',1)[1] if ':' in pr['piece'] else ''
    s=sh.get(sid,{})
    for c in pr['centers']:
        ps=c.get('personnel') or []
        m,o,dd=main(ps)
        row=['USA' if pr['country']=='US' else 'Canada', pr['region'].split('-',1)[-1], c.get('city',''), c['name'], c.get('type',''),
             s.get('name',''), head(s.get('headShliach')), m, o, dd, len(ps), round(lat,5), round(lon,5),
             (f"https://www.chabad.org/jewish-centers/{c.get('pageId') or c['id']}/{c.get('slug','')}" if (c.get('pageId') or c['id']).isdigit() else ''), c['id']]
        C.append(row)
        for p in ps:
            P.append([row[0],row[1],row[2],c['name'],row[5],p.get('title') or '',p.get('firstName') or '',p.get('lastName') or '',p.get('position') or '','Yes' if p.get('isDirector') else '','Yes' if p.get('isDeceased') else '',c['id']])
C.sort(key=lambda r:(r[0]!='USA',r[1],r[2],r[3])); P.sort(key=lambda r:(r[0]!='USA',r[1],r[2],r[3]))
wb=Workbook()
def sheet(ws,hdr,rows,widths):
    ws.append(hdr)
    for r in rows: ws.append(r)
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows():
        for cell in row: cell.font=Font(name='Arial',size=10,bold=cell.row==1,color='FFFFFF' if cell.row==1 else None)
    for cell in ws[1]: cell.fill=PatternFill('solid',fgColor='1F3A5F'); cell.alignment=Alignment(vertical='center')
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
ws=wb.active; ws.title='Centers'
sheet(ws,['Country','State/Province','City','Center','Type','Shetach','Head shliach (shetach)','Main shliach','Other shluchim','Deceased','# people listed','Latitude','Longitude','chabad.org page','Center ID'],C,[9,8,18,42,20,24,22,24,40,28,8,10,10,20,10])
for r in range(2,ws.max_row+1):
    c=ws.cell(r,14)
    if c.value: c.hyperlink=c.value; c.value='Open'; c.font=Font(name='Arial',size=10,color='0563C1',underline='single')
sheet(wb.create_sheet('Shluchim'),['Country','State/Province','City','Center','Shetach','Title','First name','Last name','Position','Director','Deceased','Center ID'],P,[9,8,18,42,24,10,16,18,22,8,9,10])
n=wb.create_sheet('About')
for line in ['Chabad centers in the USA and Canada, with their shluchim.',f'{len(C)} centers, {len(P)} people listed.','Sources: chabad.org center directory and personnel listings, as used by the Shetachim map.','Main shliach: the first living man listed. If none and a man is listed obm: "<last name> family". If only women: the first woman.','Other shluchim: everyone else alive. Deceased: those listed obm.','Shluchim sheet: one row per person, exactly as chabad.org lists them.']: n.append([line])
n.column_dimensions['A'].width=110
for r in n.iter_rows():
    for c in r: c.font=Font(name='Arial',size=10)
wb.save('exports/us-canada-centers-shluchim.xlsx')
import csv
with open('exports/us-canada-centers-shluchim.csv','w',newline='') as fh:
    w=csv.writer(fh); w.writerow([c.value for c in ws[1]][:13]+['chabad.org page','Center ID']); w.writerows(C)
print(len(C),len(P),sum(1 for r in C if r[7]), sum(1 for r in C if not r[13]))
